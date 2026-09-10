"""
Guards on the progress document.

The document exists to be handed to a manager, so the failure that matters is not
that it crashes — it is that it quietly carries a number the pipeline no longer
produces. These tests read the finished document back and check every headline
figure against the JSON on disk, so a stale figure fails the build rather than
reaching a meeting.
"""

from __future__ import annotations

import json
import re

import pytest

pytest.importorskip("docx")

from docx import Document  # noqa: E402

import cli  # noqa: E402
from contracts import schemas as S  # noqa: E402
from contracts.config import RunConfig  # noqa: E402
from docs.progress.make_progress_docx import build_document  # noqa: E402


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    data_dir = tmp_path_factory.mktemp("progress-data")
    for cmd in ("build", "run", "score"):
        assert cli.main([cmd, "--preset", "toy", "--data-dir", str(data_dir)]) == 0

    cfg = RunConfig(preset="toy", data_dir=data_dir)
    out = build_document(cfg, tmp_path_factory.mktemp("progress-out") / "P.docx")

    doc = Document(out)
    text = "\n".join(p.text for p in doc.paragraphs)
    text += "\n" + "\n".join(
        c.text for t in doc.tables for r in t.rows for c in r.cells
    )
    summary = json.loads((cfg.results_dir / "summary.json").read_text(encoding="utf-8"))
    score = json.loads((cfg.results_dir / "defect_score.json").read_text(encoding="utf-8"))
    return {"cfg": cfg, "path": out, "doc": doc, "text": text,
            "summary": summary, "score": score}


def _numbers(text: str) -> set[str]:
    """Every integer-ish token in the document, normalised without separators."""
    return {m.replace(",", "") for m in re.findall(r"\d[\d,]*", text)}


def test_document_builds_and_is_not_empty(built):
    assert built["path"].stat().st_size > 20_000
    assert len(built["doc"].tables) >= 8


def test_all_required_sections_are_present(built):
    headings = [
        p.text for p in built["doc"].paragraphs
        if p.style.name.startswith("Heading") or p.style.name == "Title"
    ]
    joined = " | ".join(headings)
    for expected in (
        "The problem",
        "What we set out to prove",
        "How we prove it without the client",
        "Progress log",
        "The numbers so far",
        "What happens next",
        "Questions a sceptic would ask",
        "Glossary",
    ):
        assert expected in joined, f"missing section: {expected}"


def test_headline_defect_numbers_match_the_json(built):
    """The whole point: the document cannot drift from what the pipeline produced."""
    s = built["score"]["summary"]
    found = _numbers(built["text"])
    for key in ("planted", "found", "missed", "false_alarms"):
        assert str(s[key]) in found, (
            f"{key}={s[key]} from defect_score.json does not appear in the document"
        )


def test_per_check_rows_match_the_json(built):
    text = built["text"]
    for row in built["score"]["by_type"]:
        assert str(row["planted"]) in _numbers(text)
        assert str(row["found"]) in _numbers(text)


def test_inventory_figures_match_the_json(built):
    inv = built["summary"]["inventory"]
    counts = built["summary"]["counts"]
    found = _numbers(built["text"])

    assert f"{round(inv['total_stock_value_sar']):,}".replace(",", "") in found
    assert str(counts["materials"]) in found
    assert str(counts["movements"]) in found
    assert str(inv["idle_24m_count"]) in found

    # percentages are rendered to one decimal place
    assert f"{inv['idle_24m_share'] * 100:.1f}%" in built["text"]
    assert f"{inv['dead_value_share'] * 100:.1f}%" in built["text"]


def test_run_date_comes_from_the_manifest(built):
    manifest = json.loads(
        (built["cfg"].results_dir / S.RUN_MANIFEST_FILE).read_text(encoding="utf-8")
    )
    assert manifest["generated_at"][:10] in built["text"]
    assert "Last updated" in built["text"]


def test_cutoff_date_is_explained(built):
    manifest = json.loads(
        (built["cfg"].results_dir / S.RUN_MANIFEST_FILE).read_text(encoding="utf-8")
    )
    assert manifest["cutoff_date"] in built["text"]


def test_no_jargon_leaks_into_the_prose(built):
    """
    Audience is a manager with no software background. File names, code identifiers
    and internal constants have no business in the document.
    """
    text = built["text"]
    for banned in (
        ".parquet", ".json", ".py", "cfg.", "results/", "engine/", "generator/",
        "DataFrame", "NEGATIVE_STOCK", "BLANK_MPN", "PROFILE_TO_SBC", "rapidfuzz",
        "TF-IDF", "Croston", "pandas",
    ):
        assert banned not in text, f"jargon leaked into the document: {banned!r}"


def test_stays_within_fifteen_pages(built):
    """
    No page count without rendering, so this bounds the content instead: roughly 450
    words to a page, plus tables and the screenshot.
    """
    words = len(built["text"].split())
    assert words < 6000, f"{words} words is heading past fifteen pages"


def test_missing_later_stages_are_absent_not_invented(built):
    """
    Stages that have not run must leave no trace. A document that carries a forecast
    heading with an empty table reads as a broken system rather than an unfinished
    one.
    """
    text = built["text"]
    assert "Stage 1" in text and "Stage 2" in text
    if "classifier_report.json" not in [
        p.name for p in built["cfg"].results_dir.iterdir()
    ]:
        assert "Stage 3" not in text
        assert "Still to come in this table" in text


def test_glossary_defines_the_terms_the_document_uses(built):
    text = built["text"]
    for term in ("Answer key", "Criticality", "Dead stock", "Lead time",
                 "Stockout", "Work order", "Safety stock"):
        assert term in text, f"glossary is missing {term}"
